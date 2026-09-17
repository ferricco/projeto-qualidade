import pandas as pd
import sqlite3
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference

def formatar_cabecalho (celula):
    celula.font = Font(bold=True, color="FFFFFF")
    celula.fill = PatternFill("solid", fgColor="1F4E79")
    celula.alignment = Alignment(horizontal="center")

def formatar_titulo (celula, texto):
    celula.value = texto
    celula.font = Font(bold=True, color="1F4E79")
    celula.alignment = Alignment(horizontal="center")

conexao = sqlite3.connect("qualidade.db")

df_nc = pd.read_sql("SELECT * FROM nao_conformidades", conexao)
df_rec = pd.read_sql("SELECT * FROM reclamacoes_clientes", conexao)

total_nc = len(df_nc)
total_rec = len(df_rec)

status_nc = df_nc["status"].value_counts()
setor_nc = df_nc["setor"].value_counts()

status_rec = df_rec["status"].value_counts()
cliente_rec = df_rec["cliente"].value_counts()

print("Total de NCs:", total_nc)
print("Total de reclamações:", total_rec)
print("\nNCs por Status:\n", status_nc)
print("\nNCs por Setor:\n", setor_nc)
print("\nStatus das reclamações:\n", status_rec)
print("\nN° de Clientes que reclamaram:\n", cliente_rec)

writer = pd.ExcelWriter("relatorio_semanal.xlsx",engine="openpyxl")
resumo = pd.DataFrame({"Indicador": ["Total de NCs", "Total de Reclamações"],
                       "Quantidade": [total_nc, total_rec]})

df_status_nc = status_nc.reset_index()
df_status_nc.columns = ["Status", "Quantidade"]

df_setor_nc = setor_nc.reset_index()
df_setor_nc.columns = ["Setores", "Quantidade"]

df_status_rec = status_rec.reset_index()
df_status_rec.columns = ["Status", "Quantidade"]

df_cliente_rec = cliente_rec.reset_index()
df_cliente_rec.columns = ["Clientes", "Qtd Reclamações"]

resumo.to_excel(writer, sheet_name="resumo", index=False, startrow=1)

df_status_nc.to_excel(writer, sheet_name="nao_conformidades", index=False, startrow=1)
df_setor_nc.to_excel(writer, sheet_name="nao_conformidades", index=False, startrow=10)

df_status_rec.to_excel(writer, sheet_name="reclamacoes", index=False,startrow=1, startcol=0)
df_cliente_rec.to_excel(writer, sheet_name="reclamacoes", index=False, startrow=1, startcol=3)

writer.close()

wb = load_workbook("relatorio_semanal.xlsx")

ws_resumo = wb["resumo"]
ws_resumo.column_dimensions["A"].width = 30
ws_resumo.column_dimensions["B"].width = 15

ws_resumo.merge_cells("A1:B1")
formatar_titulo(ws_resumo["A1"], "Resumo Total")

formatar_cabecalho(ws_resumo["A2"])
formatar_cabecalho(ws_resumo["B2"])

ws_nc = wb["nao_conformidades"]
ws_nc.column_dimensions["A"].width = 22
ws_nc.column_dimensions["B"].width = 22

ws_nc.merge_cells("A1:B1")
formatar_titulo(ws_nc["A1"], "Status de NCs")
ws_nc.merge_cells("A10:B10")
formatar_titulo(ws_nc["A10"], "NCs por Setores")

formatar_cabecalho(ws_nc["A2"])
formatar_cabecalho(ws_nc["B2"])
formatar_cabecalho(ws_nc["A11"])
formatar_cabecalho(ws_nc["B11"])

ws_rec = wb["reclamacoes"]
ws_rec.column_dimensions["A"].width = 22
ws_rec.column_dimensions["B"].width = 22
ws_rec.column_dimensions["D"].width = 25
ws_rec.column_dimensions["E"].width = 25

ws_rec.merge_cells("A1:B1")
formatar_titulo(ws_rec["A1"], "Status das Reclamações")
ws_rec.merge_cells("D1:E1")
formatar_titulo(ws_rec["D1"], "Reclamações de Clientes")

formatar_cabecalho(ws_rec["A2"])
formatar_cabecalho(ws_rec["B2"])
formatar_cabecalho(ws_rec["D2"])
formatar_cabecalho(ws_rec["E2"])

grafico_nc = BarChart()
grafico_nc.title = "NC por Status"
grafico_nc.style = 10
grafico_nc.y_axis.title = "Quantidade"
grafico_nc.x_axis.title = "Status"

dados_nc = Reference(ws_nc, min_col=2, max_col=2, min_row=2, max_row=6)
categorias_nc = Reference(ws_nc, min_col=1, max_col=1, min_row=3, max_row=6)

grafico_nc.add_data(dados_nc, titles_from_data=True)
grafico_nc.set_categories(categorias_nc)

grafico_nc.shape = 10
ws_nc.add_chart(grafico_nc, "D2")

grafico_rec = BarChart()
grafico_rec.title = "Reclamações por cliente"
grafico_rec.style = 10
grafico_rec.y_axis.title = "Quantidade"
grafico_rec.x_axis.title = "Clientes"

dados_rec = Reference(ws_rec, min_col=5, max_col=5, min_row=2, max_row=7)
clientes_rec = Reference(ws_rec, min_col=4, max_col=4, min_row=3, max_row=7)

grafico_rec.add_data(dados_rec, titles_from_data=True)
grafico_rec.set_categories(clientes_rec)

grafico_rec.shape = 10
ws_rec.add_chart(grafico_rec, "G2")


wb.save("relatorio_semanal.xlsx")

print("\nRelatório gerado com sucesso!")
